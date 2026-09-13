"""Acceptance tests for deterministic assurance exports."""

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.document import Document as WordDocument

from sodif.demo.models import FlightReport
from sodif.demo.runner import run_default_flight, run_transversal_memory_flight
from sodif.reporting import (
    LEDGER_FILENAME,
    RECEIPT_FILENAME,
    RunLedgerIntegrityError,
    persist_flight_run,
    verify_run_ledger,
)
from sodif.reporting.cli import main as export_main
from sodif.reporting.files import write_exports
from sodif.reporting.serializers import serialize_report
from sodif.reporting.service import build_flight_exports


def test_export_package_is_deterministic_and_self_verifying() -> None:
    report = run_default_flight()
    first = build_flight_exports(report)
    second = build_flight_exports(report)

    assert first == second
    with ZipFile(BytesIO(first.bundle.data)) as archive:
        names = set(archive.namelist())
        manifest = json.loads(archive.read(first.manifest.filename))
        assert names == {
            first.document.filename,
            first.report.filename,
            first.audit_log.filename,
            first.manifest.filename,
        }
        for record in manifest["files"]:
            artifact = archive.read(record["path"])
            assert record["bytes"] == len(artifact)
            assert record["digest"].startswith("sha256:")
    assert manifest["protocol"] == "sodif.evidence-manifest/v1"
    assert manifest["report_id"] == report.report_id
    assert manifest["flight_kind"] == "security"
    assert manifest["organization"] == "TECHSUITE SRL"
    assert manifest["session_id"] == "session-local-default"
    assert manifest["evidence_root"].startswith("sha256:")


def test_structured_report_and_audit_log_preserve_domain_evidence() -> None:
    report = run_default_flight()
    exports = build_flight_exports(report)
    restored = FlightReport.model_validate_json(exports.report.data)
    events = [json.loads(line) for line in exports.audit_log.data.decode().splitlines()]

    assert serialize_report(restored) == exports.report.data
    assert events[0]["event_type"] == "flight.opened"
    assert events[0]["organization"] == "TECHSUITE SRL"
    assert events[0]["route_id"] == "erp.purchase-orders"
    assert events[-1]["event_type"] == "flight.sealed"
    decisions = [event for event in events if event["event_type"] == "scenario.decision"]
    assert {event["scenario"] for event in decisions} == {
        result.scenario_id.value for result in report.results
    }
    assert sum("archive_id" in event for event in decisions) == 0

    transversal_events = [
        json.loads(line)
        for line in build_flight_exports(run_transversal_memory_flight())
        .audit_log.data.decode()
        .splitlines()
    ]
    transversal_decisions = [
        event for event in transversal_events if event["event_type"] == "scenario.decision"
    ]
    gateway_decisions = [
        event for event in transversal_events if event["event_type"] == "gateway.decision"
    ]
    assert sum("archive_id" in event for event in transversal_decisions) == 5
    assert sum(len(event.get("archive_ids", [])) for event in transversal_decisions) == 6
    assert len(gateway_decisions) == 5
    assert sum(event["status"] == "routed" for event in gateway_decisions) == 3
    assert sum(event["status"] == "blocked" for event in gateway_decisions) == 2
    assert all(event["checks"] for event in gateway_decisions)
    assert all(event["route_id"] == "erp.purchase-orders" for event in gateway_decisions)
    assert sum("gateway_decision_ids" in event for event in transversal_decisions) == 4


def test_word_report_contains_the_decision_register_and_scenario_evidence() -> None:
    exports = build_flight_exports(run_default_flight())
    document = Document(BytesIO(exports.document.data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "RAPORT SECURITY FLIGHT" in text
    assert "Configurația validată în laborator" in text
    assert "Semnătură Ed25519 detașată" in _table_text(document)
    assert "Registrul deciziilor" in text
    assert "0.14.0-dms5" not in text
    assert "Comandă autentică și neambiguă" in text
    assert "Arhivă: arc-" not in text
    assert "Document modificat după semnare" in text
    assert "Permis prezentat din nou" in text
    assert document.tables[1].rows[0].cells[0].text == "Situație"
    assert len(document.tables[1].rows) == 7
    assert len(document.tables) == 2


def test_transversal_word_report_exposes_archive_and_gateway_evidence() -> None:
    exports = build_flight_exports(run_transversal_memory_flight())
    document = Document(BytesIO(exports.document.data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "RAPORT TRANSVERSAL FLIGHT" in text
    assert "Arhivă: arc-" in text
    assert "Decizie Gateway: gateway-" in text
    assert "Politică rută: erp.purchase-orders" in text
    assert "Organizație: TECHSUITE SRL" in text
    assert "SODIF Security" in _table_text(document)
    assert "SODIF Archive" in _table_text(document)
    assert "SODIF Gateway" in _table_text(document)
    assert "Sinteza modulelor și a dovezilor" in text
    assert "Evidența arhivei documentare verificabile" in text
    assert "Decizii de rutare și blocare" in text
    assert "Trasabilitatea end-to-end între module" in text
    assert "Exemplu de corelare completă" in text
    assert "Document modificat după semnare" in text
    assert "Neînregistrat" in "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert len(document.tables) == 6


def test_cli_writer_persists_every_artifact(tmp_path: Path) -> None:
    exports = build_flight_exports(run_default_flight())
    paths = write_exports(tmp_path, exports)

    assert {path.name for path in paths} == {
        exports.document.filename,
        exports.report.filename,
        exports.audit_log.filename,
        exports.manifest.filename,
        exports.bundle.filename,
    }
    assert all(path.is_file() and path.stat().st_size > 0 for path in paths)


def test_cli_exports_both_product_flights(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    output = tmp_path / "exports"
    archive = tmp_path / "archive"

    export_main(
        [
            "--kind",
            "both",
            "--output-dir",
            str(output),
            "--archive-root",
            str(archive),
        ]
    )

    summary = json.loads(capsys.readouterr().out)
    assert summary["status"] == "exported"
    assert [flight["flight_kind"] for flight in summary["flights"]] == [
        "security",
        "transversal",
    ]
    assert len(summary["artifacts"]) == 10
    assert len(tuple(output.glob("*.docx"))) == 2


def test_completed_run_is_persisted_by_session_kind_and_report(tmp_path: Path) -> None:
    report = run_transversal_memory_flight()
    exports = build_flight_exports(report)

    persisted = persist_flight_run(tmp_path, report, exports)

    assert persisted.directory == (
        tmp_path / report.configuration.session_id / "transversal" / report.report_id
    )
    assert {path.name for path in persisted.artifacts} == {
        exports.document.filename,
        exports.report.filename,
        exports.audit_log.filename,
        exports.manifest.filename,
        exports.bundle.filename,
        RECEIPT_FILENAME,
    }
    assert all(path.is_file() and path.stat().st_size > 0 for path in persisted.artifacts)
    assert persisted.ledger_path == tmp_path / LEDGER_FILENAME
    assert persisted.ledger_verification.valid
    assert persisted.ledger_verification.entries == 1


def test_run_integrity_ledger_chains_both_flights_and_issues_receipts(tmp_path: Path) -> None:
    security_report = run_default_flight()
    transversal_report = run_transversal_memory_flight()

    first = persist_flight_run(
        tmp_path,
        security_report,
        build_flight_exports(security_report),
    )
    second = persist_flight_run(
        tmp_path,
        transversal_report,
        build_flight_exports(transversal_report),
    )
    verification = verify_run_ledger(tmp_path)
    receipt = json.loads(second.ledger_receipt.read_bytes())

    assert verification.valid
    assert verification.entries == 2
    assert verification.head_digest == second.ledger_entry.entry_digest
    assert first.ledger_entry.sequence == 1
    assert first.ledger_entry.previous_entry_digest is None
    assert second.ledger_entry.sequence == 2
    assert second.ledger_entry.previous_entry_digest == first.ledger_entry.entry_digest
    assert receipt["protocol"] == "sodif.run-integrity-receipt/v1"
    assert receipt["entry_digest"] == second.ledger_entry.entry_digest
    assert receipt["verified_entries"] == 2
    assert receipt["verified_head_digest"] == verification.head_digest


def test_run_integrity_ledger_detects_tampering_and_refuses_a_new_entry(
    tmp_path: Path,
) -> None:
    report = run_default_flight()
    exports = build_flight_exports(report)
    persisted = persist_flight_run(tmp_path, report, exports)
    ledger = persisted.ledger_path
    original = ledger.read_text(encoding="utf-8")
    ledger.write_text(
        original.replace(
            '"organization":"TECHSUITE SRL"',
            '"organization":"TECHSUITX SRL"',
            1,
        ),
        encoding="utf-8",
    )

    verification = verify_run_ledger(tmp_path)

    assert not verification.valid
    assert "entry digest does not match its content" in " ".join(verification.errors)
    with pytest.raises(RunLedgerIntegrityError, match="verification failed"):
        persist_flight_run(tmp_path, report, exports)


def test_run_integrity_ledger_detects_a_modified_audit_package(tmp_path: Path) -> None:
    report = run_default_flight()
    exports = build_flight_exports(report)
    persisted = persist_flight_run(tmp_path, report, exports)
    bundle = persisted.directory / exports.bundle.filename
    bundle.write_bytes(bundle.read_bytes() + b"tampered")

    verification = verify_run_ledger(tmp_path)

    assert not verification.valid
    assert "audit package digest does not match" in " ".join(verification.errors)


def _table_text(document: WordDocument) -> str:
    return "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
