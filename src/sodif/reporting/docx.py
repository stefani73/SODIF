"""Professional deterministic DOCX assurance report."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from sodif.demo.models import FlightKind, FlightReport
from sodif.domain.enums import DocumentSecurityMode
from sodif.ui.presentation import FlightView, ScenarioView, present_flight

_NAVY = "0B2545"
_BLUE = "2E74B5"
_DARK_BLUE = "1F4D78"
_TEAL = "087F8C"
_MUTED = "607086"
_LIGHT_FILL = "F2F4F7"
_SUCCESS_FILL = "EAF7F3"
_SUCCESS = "087A64"
_WARNING_FILL = "FFF5DF"
_WARNING = "A8660C"
_DANGER_FILL = "FFF0F1"
_DANGER = "B33B47"
_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def render_docx_report(report: FlightReport) -> bytes:
    """Render one report using the standard-business-brief design preset."""
    document = Document()
    _configure_document(document, report)
    view = present_flight(report)
    _add_masthead(document, report, view)
    _add_lab_configuration(document, report)
    _add_decision_register(document, view)
    if report.flight_kind is FlightKind.TRANSVERSAL:
        if report.security_mode is DocumentSecurityMode.ADVANCED:
            _add_transversal_evidence(document, report)
        else:
            _add_standard_transfer_evidence(document, report)
    compact_scenarios = report.flight_kind is FlightKind.TRANSVERSAL
    for scenario in view.scenarios:
        _add_scenario(document, scenario, compact=compact_scenarios)
    buffer = BytesIO()
    document.save(buffer)
    return _normalize_docx(buffer.getvalue())


def _configure_document(document: WordDocument, report: FlightReport) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    properties = document.core_properties
    properties.title = "Raport operațional SODIF"
    properties.subject = "Validarea în laborator a controlului execuției document-to-API"
    properties.author = "SODIF"
    properties.last_modified_by = "SODIF"
    properties.created = report.started_at.replace(tzinfo=None)
    properties.modified = report.completed_at.replace(tzinfo=None)
    properties.revision = 1
    properties.comments = ""

    _configure_styles(document)
    _configure_numbering(document)
    _configure_header(section.header.paragraphs[0])
    _configure_footer(section.footer.paragraphs[0], report.report_id)


def _configure_styles(document: WordDocument) -> None:
    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", 11, _NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, _BLUE, 16, 8),
        "Heading 2": (13, _BLUE, 12, 6),
        "Heading 3": (12, _DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        _set_style_font(style, "Calibri", size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    _set_style_font(bullet, "Calibri", 11, _NAVY)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8)
    bullet.paragraph_format.line_spacing = 1.167

    evidence = document.styles.add_style("SODIF Evidence", WD_STYLE_TYPE.PARAGRAPH)
    evidence.base_style = normal
    _set_style_font(evidence, "Consolas", 9, _DARK_BLUE)
    evidence.paragraph_format.left_indent = Inches(0.16)
    evidence.paragraph_format.space_after = Pt(4)
    evidence.paragraph_format.keep_together = True


def _set_style_font(
    style: object,
    name: str,
    size: float,
    color: str,
    *,
    bold: bool | None = None,
) -> None:
    font = style.font  # type: ignore[attr-defined]
    font.name = name
    font.size = Pt(size)
    font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        font.bold = bold
    rpr = style._element.get_or_add_rPr()  # type: ignore[attr-defined]
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def _configure_numbering(document: WordDocument) -> None:
    numbering = document.part.numbering_part.element
    for level in numbering.iter(qn("w:lvl")):
        number_format = level.find(qn("w:numFmt"))
        if number_format is None or number_format.get(qn("w:val")) != "bullet":
            continue
        properties = level.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            level.append(properties)
        indent = properties.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            properties.append(indent)
        indent.set(qn("w:left"), "720")
        indent.set(qn("w:hanging"), "360")
        tabs = properties.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            properties.insert(0, tabs)
        for child in list(tabs):
            tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)


def _configure_header(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("SODIF  |  SECURITATEA EXECUȚIEI")
    _set_run_font(run, "Calibri", 8, _MUTED, bold=True)


def _configure_footer(paragraph: Paragraph, report_id: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(f"{report_id}  |  Page ")
    _set_run_font(run, "Calibri", 8, _MUTED)
    field_run = paragraph.add_run()
    _set_run_font(field_run, "Calibri", 8, _MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, separate, field_text, end))


def _add_masthead(document: WordDocument, report: FlightReport, view: FlightView) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(14)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(6)
    _set_run_font(kicker.add_run("CYBERSECURITY EXECUTION CONTROL"), "Calibri", 9, _TEAL, bold=True)

    flight_title = {
        FlightKind.SECURITY: "RAPORT SECURITY FLIGHT",
        FlightKind.TRANSVERSAL: "RAPORT TRANSVERSAL FLIGHT",
    }[report.flight_kind]
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    _set_run_font(title.add_run(flight_title), "Calibri", 24, _NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True
    subtitle_text = (
        "Validarea lanțului de securitate de la revizia semnată la efectul API"
        if report.security_mode is DocumentSecurityMode.ADVANCED
        else "Validarea traseului standard de la revizia semnată la transferul API direct"
    )
    _set_run_font(
        subtitle.add_run(subtitle_text),
        "Calibri",
        12,
        _MUTED,
    )

    coverage = (
        "SODIF Security / API controlat"
        if report.flight_kind is FlightKind.SECURITY
        else (
            "SODIF Security / SODIF Archive / SODIF Gateway / API controlat"
            if report.security_mode is DocumentSecurityMode.ADVANCED
            else "Semnătură / SODIF Archive / adaptor API direct"
        )
    )
    security_label = (
        "Protecție avansată SODIF"
        if report.security_mode is DocumentSecurityMode.ADVANCED
        else "Transfer standard direct"
    )
    metadata = (
        ("Organizație", report.configuration.organization_name),
        ("Spațiu operațional", report.configuration.workspace_name),
        ("Domeniu", report.configuration.domain_name),
        ("Mediu", report.configuration.environment),
        ("Raport", report.report_id),
        ("Rezultat", "CONFORM" if report.passed else "NECONFORM"),
        ("Regim", security_label),
        ("Sigilat la", report.completed_at.isoformat().replace("+00:00", "Z")),
        ("Acoperire", coverage),
        (
            "Serviciu protejat",
            f"{report.configuration.protected_service} · {report.configuration.route_id}",
        ),
        ("Versiune", report.release),
    )
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        _set_run_font(paragraph.add_run(f"{label}: "), "Calibri", 10, _NAVY, bold=True)
        _set_run_font(paragraph.add_run(value), "Calibri", 10, _NAVY)

    document.add_paragraph().paragraph_format.space_after = Pt(4)
    _add_callout(document, view.title, view.detail, "success" if report.passed else "danger")


def _add_lab_configuration(document: WordDocument, report: FlightReport) -> None:
    """Describe the reproducible laboratory boundary represented by the report."""
    document.add_heading("Configurația validată în laborator", level=1)
    introduction = document.add_paragraph(
        "Rularea exercită implementarea modulară Python prin date controlate, chei de test "
        "și un serviciu API local. Rezultatele sunt reproductibile și corelate cu artefactele "
        "exportate."
    )
    introduction.paragraph_format.keep_with_next = True
    rows = [
        (
            "Revizie semnată",
            "Semnătură Ed25519 detașată peste metadate canonice și digestul SHA-256 al "
            "conținutului.",
        )
    ]
    if report.security_mode is DocumentSecurityMode.ADVANCED:
        rows.extend(
            (
                (
                    "Reprezentări confruntate",
                    "Citire structurală cu pypdf și două trasee vizuale distincte, randate "
                    "prin MuPDF și Poppler și citite cu același motor Tesseract.",
                ),
                (
                    "Verificare adaptivă",
                    "Niveluri V0-V3 selectate în funcție de risc, completitudinea valorilor "
                    "și divergențele cu efect operațional.",
                ),
                (
                    "Dovadă și permis",
                    "Angajamente SHA-256 pe câmp, rădăcină Merkle și permis Ed25519 cu "
                    "utilizare unică, legate de acțiunea API.",
                ),
                (
                    "Control API",
                    "Motor Gateway local pentru politică, compararea amprentelor și blocarea "
                    "reutilizării permisului.",
                ),
            )
        )
    else:
        rows.extend(
            (
                (
                    "Date de integrare",
                    "Valorile configurate sunt mapate determinist în corpul cererii API.",
                ),
                (
                    "Transfer direct",
                    "Cererea este transmisă adaptorului API local fără consens semantic, "
                    "dovadă de invariabilitate, permis sau control Gateway.",
                ),
            )
        )
    if report.flight_kind is FlightKind.TRANSVERSAL:
        rows.append(
            (
                "Arhivă documentară",
                "Index SQLite, obiecte locale, verificarea digestului la citire și "
                "continuitatea criptografică a reviziilor.",
            )
        )
    rows.append(
        (
            "Audit",
            "Raport DOCX, date JSON, jurnal NDJSON, manifest SHA-256, pachet ZIP și "
            "registru criptografic al rulărilor.",
        )
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    widths = (2450, 6910)
    _set_table_geometry(table, widths)
    headers = ("Componentă", "Implementare validată")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9.2)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for label, value in rows:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], label, bold=True, size=9)
        _set_cell_text(cells[1], value, size=9)


def _add_callout(
    document: WordDocument,
    title: str,
    detail: str,
    tone: str,
    *,
    compact: bool = False,
) -> None:
    fill, accent = {
        "success": (_SUCCESS_FILL, _SUCCESS),
        "warning": (_WARNING_FILL, _WARNING),
        "danger": (_DANGER_FILL, _DANGER),
    }[tone]
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3 if compact else 5)
    paragraph.paragraph_format.space_after = Pt(6 if compact else 10)
    paragraph.paragraph_format.keep_together = True
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), accent)
    borders.append(left)
    properties.append(borders)
    _set_run_font(
        paragraph.add_run(f"{title}\n"),
        "Calibri",
        10.5 if compact else 11,
        accent,
        bold=True,
    )
    _set_run_font(
        paragraph.add_run(detail),
        "Calibri",
        10 if compact else 10.5,
        _NAVY,
    )


def _add_decision_register(document: WordDocument, view: FlightView) -> None:
    document.add_heading("Registrul deciziilor", level=1)
    paragraph = document.add_paragraph(
        "Registrul sintetizează decizia observabilă și comportamentul rezultat al API-ului."
    )
    paragraph.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = (3450, 1950, 3960)
    _set_table_geometry(table, widths)
    headers = ("Situație", "Decizie", "Efect asupra API")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9.5)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for scenario in view.scenarios:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], scenario.title, size=9.5)
        _set_cell_text(
            cells[1],
            scenario.verdict,
            bold=True,
            color=_tone_color(scenario.tone),
            size=9.5,
        )
        _set_cell_text(cells[2], scenario.api_effect, size=9.5)
    note = document.add_paragraph(
        "Pachetul de audit include reprezentarea structurată a raportului, jurnalul ordonat "
        "și manifestul de integritate care permit verificarea acestor concluzii."
    )
    note.paragraph_format.space_before = Pt(8)


def _add_standard_transfer_evidence(document: WordDocument, report: FlightReport) -> None:
    """Document the intentionally direct execution path selected for the run."""
    result = report.results[0]
    if result.receipt is None:
        raise AssertionError("standard transversal flight lacks an API receipt")
    _add_chapter_heading(
        document,
        "REGIM STANDARD",
        "Traseul direct către adaptorul API",
        "Rularea separă explicit controalele de bază de protecția avansată SODIF. "
        "Documentul a fost validat criptografic și arhivat, iar valorile configurate au "
        "fost transferate direct către adaptorul API local.",
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (1850, 3250, 2300, 1960)
    _set_table_geometry(table, widths)
    for cell, text in zip(
        table.rows[0].cells,
        ("Componentă", "Operațiune", "Dovadă", "Rezultat"),
        strict=True,
    ):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    rows = (
        (
            "Semnătură",
            "Validarea reviziei și a digestului conținutului.",
            result.workflow.correlation_id,
            "VALIDAT",
        ),
        (
            "SODIF Archive",
            "Înregistrarea reviziilor în regim standard.",
            f"{len(result.archive_ids)} revizii",
            "ARHIVAT",
        ),
        (
            "Adaptor API",
            "Transferul direct al parametrilor configurați.",
            f"{result.receipt.execution_id} · HTTP {result.receipt.response_code}",
            "TRANSFERAT",
        ),
    )
    for component, operation, evidence, outcome in rows:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], component, bold=True, size=9.2)
        _set_cell_text(cells[1], operation, size=9.2)
        _set_cell_text(cells[2], evidence, size=8.9)
        _set_cell_text(cells[3], outcome, bold=True, color=_SUCCESS, size=9.2)

    document.add_heading("Delimitarea traseului", level=2)
    for item in (
        "Consensul semantic și verificarea adaptivă nu au fost executate.",
        "Nu au fost generate dovada de invariabilitate sau permisul criptografic.",
        "SODIF Gateway nu a intervenit în rutare; adaptorul API a primit cererea direct.",
        "Jurnalul și raportul identifică regimul standard pentru întreaga rulare.",
    ):
        document.add_paragraph(item, style="List Bullet")


def _add_transversal_evidence(document: WordDocument, report: FlightReport) -> None:
    """Expose the additional DMS and Gateway evidence carried by a transversal flight."""
    archived_results = tuple(result for result in report.results if result.archive_ids)
    archive_records = sum(len(result.archive_ids) for result in archived_results)
    gateway_decisions = tuple(
        decision for result in report.results for decision in result.gateway_decisions
    )
    routed = sum(decision.status.value == "routed" for decision in gateway_decisions)
    blocked = sum(decision.status.value == "blocked" for decision in gateway_decisions)

    _add_chapter_heading(
        document,
        "ACOPERIRE TRANSVERSALĂ",
        "Sinteza modulelor și a dovezilor",
        "Raportul transversal separă rezultatele fiecărui modul și păstrează legătura "
        "dintre documentul semnat, înregistrarea documentară, decizia Gateway și efectul API.",
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (1850, 3250, 2300, 1960)
    _set_table_geometry(table, widths)
    headers = ("Modul", "Responsabilitate demonstrată", "Dovadă observată", "Rezultat")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    module_rows = (
        (
            "SODIF Security",
            "Confruntă structura PDF cu forma vizibilă, dovedește valorile stabile și emite "
            "permisul unic legat de execuție.",
            f"{report.passed_scenarios}/{len(report.results)} scenarii conforme",
            "CONTROLAT",
        ),
        (
            "SODIF Archive",
            "Înregistrează reviziile validate și păstrează continuitatea versiunilor.",
            f"{archive_records} referințe de arhivare în {len(archived_results)} scenarii",
            "VERIFICABIL",
        ),
        (
            "SODIF Gateway",
            "Verifică proveniența fiecărui parametru, permisul și legarea exactă de acțiunea "
            "solicitată.",
            f"{len(gateway_decisions)} decizii: {routed} rutate / {blocked} blocate",
            "APLICAT",
        ),
    )
    for module, responsibility, evidence, result in module_rows:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], module, bold=True, size=9.2)
        _set_cell_text(cells[1], responsibility, size=9.2)
        _set_cell_text(cells[2], evidence, size=9.2)
        _set_cell_text(cells[3], result, bold=True, color=_SUCCESS, size=9.2)

    document.add_heading("Flux operațional demonstrat", level=2)
    for item in (
        "Documentul este acceptat numai dacă revizia și semnătura corespund conținutului primit.",
        "Valorile confirmate sunt angajate pe câmp, compilate într-o acțiune exactă și "
        "legate de un permis unic.",
        "Arhiva păstrează reviziile validate, iar Gateway-ul decide dacă acțiunea poate fi rutată.",
        "Fiecare rezultat rămâne corelat cu identificatorii tehnici incluși în pachetul de audit.",
    ):
        document.add_paragraph(item, style="List Bullet")

    _add_dms_evidence(document, report)
    _add_gateway_evidence(document, report)
    _add_end_to_end_traceability(document, report)


def _add_dms_evidence(document: WordDocument, report: FlightReport) -> None:
    _add_chapter_heading(
        document,
        "SODIF ARCHIVE",
        "Evidența arhivei documentare verificabile",
        "Registrul de mai jos inventariază reviziile acceptate de controlul criptografic. "
        "Documentul modificat după semnare este oprit înainte de arhivare.",
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (2700, 1150, 3300, 2210)
    _set_table_geometry(table, widths)
    headers = ("Scenariu", "Revizii", "Ultima înregistrare", "Continuitate")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for result in report.results:
        if not result.archive_ids:
            continue
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], result.title, size=9.2)
        _set_cell_text(cells[1], str(len(result.archive_ids)), bold=True, size=9.2)
        _set_cell_text(cells[2], result.archive_ids[-1], size=8.8)
        continuity = "Lanț de revizii" if len(result.archive_ids) > 1 else "Revizie validată"
        _set_cell_text(cells[3], continuity, size=9.2)

    document.add_heading("Controale demonstrate", level=2)
    for item in (
        "Arhivarea este declanșată după validarea semnăturii și a digestului reviziei.",
        "O revizie ulterioară declară explicit amprenta reviziei anterioare.",
        "Identificatorul arhivei este inclus în rezultatul scenariului și în jurnalul exportat.",
        "Respingerea timpurie împiedică introducerea în arhivă a documentelor alterate.",
    ):
        document.add_paragraph(item, style="List Bullet")


def _add_gateway_evidence(document: WordDocument, report: FlightReport) -> None:
    _add_chapter_heading(
        document,
        "SODIF GATEWAY",
        "Decizii de rutare și blocare",
        "Gateway-ul aplică politica rutei și verifică permisul criptografic direct față de "
        "acțiunea observată și dovada valorilor înainte ca cererea să ajungă la serviciul "
        "protejat.",
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (2450, 2350, 1450, 3110)
    _set_table_geometry(table, widths)
    headers = ("Scenariu", "Decizie Gateway", "Rezultat", "Politică și verificări")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for result in report.results:
        for decision in result.gateway_decisions:
            cells = table.add_row().cells
            _apply_row_geometry(cells, widths)
            _set_cell_text(cells[0], result.title, size=9.1)
            _set_cell_text(cells[1], decision.decision_id, size=8.8)
            status = "RUTATĂ" if decision.status.value == "routed" else "BLOCATĂ"
            tone = _SUCCESS if decision.status.value == "routed" else _DANGER
            _set_cell_text(cells[2], status, bold=True, color=tone, size=9.1)
            _set_cell_text(
                cells[3],
                f"{decision.route_id} · {len(decision.checks)} controale · {decision.code}",
                size=8.9,
            )

    document.add_heading("Limita de execuție demonstrată", level=2)
    for item in (
        "Fiecare parametru transmis este acoperit de un angajament al câmpului aprobat și "
        "de rădăcina criptografică a setului de valori.",
        "O rută modificată după emiterea permisului este blocată înainte de serviciul protejat.",
        "Prima prezentare conformă poate fi rutată, iar reutilizarea aceluiași permis "
        "este respinsă.",
        "Deciziile păstrează digestul cererii, digestul acțiunii autorizate și rezultatul "
        "controalelor.",
    ):
        document.add_paragraph(item, style="List Bullet")


def _add_end_to_end_traceability(document: WordDocument, report: FlightReport) -> None:
    _add_chapter_heading(
        document,
        "TRASABILITATE",
        "Trasabilitatea end-to-end între module",
        "Matricea confirmă unde s-a oprit sau a continuat fiecare scenariu și separă "
        "explicit efectul documentar de decizia de securitate și de execuția API.",
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = (2450, 2450, 1700, 2760)
    _set_table_geometry(table, widths)
    headers = ("Scenariu", "Document și intenție", "Arhivă", "Gateway și API")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=9)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for result in report.results:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], result.title, size=9)
        if result.scenario_id.value == "tampered-document":
            document_state = "Blocat la integritate"
        elif result.scenario_id.value == "semantic-conflict":
            document_state = "Validat; consens neconcludent"
        else:
            document_state = "Validat; intenție confirmată"
        _set_cell_text(cells[1], document_state, size=9)
        archive_count = len(result.archive_ids)
        archive_state = (
            f"{archive_count} {'revizie' if archive_count == 1 else 'revizii'}"
            if archive_count
            else "Neînregistrat"
        )
        _set_cell_text(cells[2], archive_state, size=9)
        gateway_states = tuple(
            "Rutată" if decision.status.value == "routed" else "Blocată"
            for decision in result.gateway_decisions
        )
        if gateway_states:
            execution = " · ".join(gateway_states)
            if result.receipt is not None:
                execution = f"{execution}; API {result.receipt.response_code}"
        else:
            execution = "Gateway neapelat; API neexecutat"
        _set_cell_text(cells[3], execution, size=9)

    reference = next(
        result
        for result in report.results
        if result.archive_ids and result.gateway_decisions and result.receipt is not None
    )
    if reference.verification is None or reference.receipt is None:
        raise AssertionError("complete transversal reference lacks verification or receipt")
    document.add_heading("Exemplu de corelare completă", level=2)
    correlation_items = (
        f"Tranzacție: {reference.workflow.correlation_id}",
        f"Document: {_compact_report_digest(reference.verification.revision_digest)}",
        f"Arhivă: {', '.join(reference.archive_ids)}",
        f"Permis: {reference.permit_id}",
        f"Provocare semantică: {_compact_report_digest(reference.challenge_digest or '')}",
        f"Rădăcină valori: {_compact_report_digest(reference.field_root or '')}",
        (f"Dovadă execuție: {_compact_report_digest(reference.execution_proof_digest or '')}"),
        f"Decizie Gateway: {reference.gateway_decisions[-1].decision_id}",
        (
            f"Execuție API: {reference.receipt.execution_id} · "
            f"răspuns {reference.receipt.response_code}"
        ),
    )
    for item in correlation_items:
        paragraph = document.add_paragraph(style="SODIF Evidence")
        _set_run_font(paragraph.add_run(item), "Consolas", 9, _DARK_BLUE)


def _add_chapter_heading(
    document: WordDocument,
    kicker_text: str,
    title: str,
    introduction: str,
) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.page_break_before = True
    kicker.paragraph_format.space_after = Pt(3)
    _set_run_font(kicker.add_run(kicker_text), "Calibri", 8.5, _TEAL, bold=True)
    document.add_heading(title, level=1)
    document.add_paragraph(introduction)


def _compact_report_digest(value: str) -> str:
    return f"{value[:18]}…{value[-8:]}"


def _add_scenario(
    document: WordDocument,
    scenario: ScenarioView,
    *,
    compact: bool = False,
) -> None:
    body_size = 10 if compact else 11
    evidence_size = 8.5 if compact else 9

    kicker = document.add_paragraph()
    kicker.paragraph_format.page_break_before = True
    kicker.paragraph_format.space_after = Pt(2 if compact else 3)
    _set_run_font(kicker.add_run(scenario.kicker.upper()), "Calibri", 8.5, _TEAL, bold=True)
    title = document.add_heading(scenario.title, level=1)
    if compact:
        title.paragraph_format.space_before = Pt(10)
        title.paragraph_format.space_after = Pt(4)
    summary = document.add_paragraph(scenario.summary)
    if compact:
        summary.paragraph_format.space_after = Pt(4)
        summary.paragraph_format.line_spacing = 1
        for run in summary.runs:
            _set_run_font(run, "Calibri", body_size, _NAVY)
    _add_callout(
        document,
        f"Decizie: {scenario.verdict}",
        scenario.verdict_detail,
        scenario.tone,
        compact=compact,
    )

    heading = document.add_heading("Rezultatul controalelor", level=2)
    if compact:
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)
    for control in scenario.controls:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        if compact:
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1
        _set_run_font(
            paragraph.add_run(f"{control.name} — {control.state}. "),
            "Calibri",
            body_size,
            _tone_color(control.tone),
            bold=True,
        )
        _set_run_font(paragraph.add_run(control.detail), "Calibri", body_size, _NAVY)

    heading = document.add_heading("Rațiunea deciziei și efectul asupra API", level=2)
    if compact:
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)
    paragraph = document.add_paragraph()
    if compact:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1
    _set_run_font(paragraph.add_run("Rațiune. "), "Calibri", body_size, _NAVY, bold=True)
    _set_run_font(paragraph.add_run(scenario.optimization_note), "Calibri", body_size, _NAVY)
    paragraph = document.add_paragraph()
    if compact:
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1
    _set_run_font(
        paragraph.add_run(f"{scenario.api_effect}. "),
        "Calibri",
        body_size,
        _NAVY,
        bold=True,
    )
    _set_run_font(paragraph.add_run(scenario.api_detail), "Calibri", body_size, _NAVY)

    if scenario.comparisons:
        _add_field_comparison_table(document, scenario, compact=compact)

    heading = document.add_heading("Traseul deciziei", level=2)
    if compact:
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)
    for timeline_item in scenario.timeline:
        paragraph = document.add_paragraph(timeline_item, style="List Bullet")
        if compact:
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1
            for run in paragraph.runs:
                _set_run_font(run, "Calibri", 9.5, _NAVY)

    heading = document.add_heading("Identificatori și trasabilitate", level=2)
    if compact:
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)
    for evidence in scenario.evidence:
        paragraph = document.add_paragraph(style="SODIF Evidence")
        if compact:
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1
        _set_run_font(
            paragraph.add_run(f"{evidence.label}: "),
            "Consolas",
            evidence_size,
            _MUTED,
            bold=True,
        )
        _set_run_font(
            paragraph.add_run(evidence.value),
            "Consolas",
            evidence_size,
            _DARK_BLUE,
        )


def _add_field_comparison_table(
    document: WordDocument,
    scenario: ScenarioView,
    *,
    compact: bool,
) -> None:
    heading = document.add_heading("Valorile confruntate", level=2)
    if compact:
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(3)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = (2100, 4160, 3100)
    _set_table_geometry(table, widths)
    for cell, text in zip(
        table.rows[0].cells,
        ("Câmp", "Valori observate", "Concluzie"),
        strict=True,
    ):
        _set_cell_text(cell, text, bold=True, color=_NAVY, size=8.8 if compact else 9.2)
        _set_cell_fill(cell, _LIGHT_FILL)
    _repeat_table_header(table)
    for comparison in scenario.comparisons:
        cells = table.add_row().cells
        _apply_row_geometry(cells, widths)
        _set_cell_text(cells[0], comparison.label, bold=True, size=8.7 if compact else 9.2)
        _set_cell_text(
            cells[1],
            "\n".join(comparison.observations),
            size=8.5 if compact else 9,
        )
        _set_cell_text(
            cells[2],
            comparison.conclusion,
            bold=True,
            color=_tone_color(comparison.tone),
            size=8.5 if compact else 9,
        )


def _set_run_font(
    run: Run,
    name: str,
    size: float,
    color: str,
    *,
    bold: bool | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def _tone_color(tone: str) -> str:
    return {
        "success": _SUCCESS,
        "warning": _WARNING,
        "danger": _DANGER,
        "neutral": _MUTED,
    }[tone]


def _set_table_geometry(table: Table, widths: tuple[int, ...]) -> None:
    if sum(widths) != 9360:
        raise ValueError("report table widths must total 9360 DXA")
    table.autofit = False
    properties = table._tbl.tblPr
    _replace_sized_element(properties, "w:tblW", "9360", type_value="dxa")
    _replace_sized_element(properties, "w:tblInd", "120", type_value="dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        _apply_row_geometry(row.cells, widths)


def _apply_row_geometry(cells: tuple[_Cell, ...], widths: tuple[int, ...]) -> None:
    for cell, width in zip(cells, widths, strict=True):
        cell.width = Inches(width / 1440)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        properties = cell._tc.get_or_add_tcPr()
        tc_width = properties.get_or_add_tcW()
        tc_width.set(qn("w:type"), "dxa")
        tc_width.set(qn("w:w"), str(width))
        margins = properties.find(qn("w:tcMar"))
        if margins is None:
            margins = OxmlElement("w:tcMar")
            properties.append(margins)
        for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
            item = margins.find(qn(f"w:{edge}"))
            if item is None:
                item = OxmlElement(f"w:{edge}")
                margins.append(item)
            item.set(qn("w:w"), str(value))
            item.set(qn("w:type"), "dxa")


def _replace_sized_element(
    parent: object,
    tag: str,
    value: str,
    *,
    type_value: str,
) -> None:
    element = parent.find(qn(tag))  # type: ignore[attr-defined]
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)  # type: ignore[attr-defined]
    element.set(qn("w:w"), value)
    element.set(qn("w:type"), type_value)


def _set_cell_text(
    cell: _Cell,
    text: str,
    *,
    bold: bool = False,
    color: str = _NAVY,
    size: float = 10,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    _set_run_font(paragraph.add_run(text), "Calibri", size, color, bold=bold)


def _set_cell_fill(cell: _Cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_table_header(table: Table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _normalize_docx(data: bytes) -> bytes:
    source = BytesIO(data)
    output = BytesIO()
    with (
        ZipFile(source, "r") as reader,
        ZipFile(
            output,
            "w",
            compression=ZIP_DEFLATED,
            compresslevel=9,
        ) as writer,
    ):
        for original in sorted(reader.infolist(), key=lambda item: item.filename):
            if original.is_dir():
                continue
            info = ZipInfo(original.filename, date_time=_ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0o600 << 16
            writer.writestr(
                info,
                reader.read(original.filename),
                compress_type=ZIP_DEFLATED,
                compresslevel=9,
            )
    return output.getvalue()
